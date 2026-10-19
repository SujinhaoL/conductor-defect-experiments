"""Append/update only the current FormalBaseline row in the existing workbook.

No workbook is rebuilt. Every mutation backs up XLSX and CSV, stages both files,
checks historical cells/styles, and then replaces the pair (rollback on error).
"""
from __future__ import annotations

from contextlib import contextmanager
from copy import copy
import csv
from datetime import date, datetime
import hashlib
import json
import os
from pathlib import Path
import shutil
import uuid

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter


def sha256(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


@contextmanager
def exclusive_lock(path: Path, owner: dict):
    try:
        stream = path.open('x', encoding='utf-8')
    except FileExistsError as exc:
        raise RuntimeError(f'锁文件已存在，拒绝并发或自动重跑：{path}；请先检查对应运行状态。') from exc
    try:
        with stream:
            json.dump({'pid': os.getpid(), **owner}, stream, ensure_ascii=False, indent=2)
        yield
    finally:
        path.unlink()  # Only this process's newly-created lock, never a pre-existing one.


def read_registry(path: Path):
    wb = load_workbook(path)
    try:
        if wb.sheetnames != ['Runs', 'ProjectSetup', 'FieldDefinitions']:
            raise RuntimeError(f'登记表工作表结构发生变化：{wb.sheetnames}')
        ws = wb['Runs']
        candidates = [r[0].row for r in ws.iter_rows() if r[0].value == 'experiment_id']
        if len(candidates) != 1:
            raise RuntimeError('无法唯一定位 Runs 列名行。')
        header_row = candidates[0]
        headers = [c.value for c in ws[header_row]]
        if None in headers or len(set(headers)) != len(headers):
            raise RuntimeError('Runs 列名缺失或重复，停止。')
        records = []
        for row in ws.iter_rows(min_row=header_row + 1):
            if row[0].value is None:
                if any(c.value is not None for c in row):
                    raise RuntimeError('Runs 存在没有 experiment_id 的非空行。')
                continue
            records.append((row[0].row, dict(zip(headers, [c.value for c in row]))))
        ids = [r['experiment_id'] for _, r in records]
        if len(set(ids)) != len(ids):
            raise RuntimeError('登记表有重复 experiment_id，停止。')
        setup = {r[0].value: r[1].value for r in wb['ProjectSetup'].iter_rows(min_row=2) if r[0].value}
        definitions = [[c.value for c in row] for row in wb['FieldDefinitions'].iter_rows()]
        schema = hashlib.sha256(json.dumps(
            {'headers': headers, 'definitions': definitions}, ensure_ascii=False, sort_keys=True
        ).encode()).hexdigest()
        return wb, header_row, headers, records, setup, schema
    except BaseException:
        wb.close()
        raise


def csv_value(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return str(value)


def csv_matches(value, text):
    if isinstance(value, bool):
        return text.lower() == str(value).lower()
    if isinstance(value, (int, float)):
        try:
            return float(text) == value
        except ValueError:
            return False
    if isinstance(value, (date, datetime)):
        return text in (value.strftime('%Y-%m-%d'), str(value), value.isoformat())
    return csv_value(value) == text


def check_csv_pair(path: Path, headers, records):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != headers:
            raise RuntimeError(f'CSV 列名与 Runs 不一致：{path}')
        rows = list(reader)
    if len(rows) != len(records):
        raise RuntimeError(f'XLSX/CSV 行数不一致：{len(records)}/{len(rows)}')
    for (_, expected), actual in zip(records, rows):
        for key in headers:
            if not csv_matches(expected[key], actual.get(key, '')):
                raise RuntimeError(f'XLSX/CSV 不一致：{expected["experiment_id"]}.{key}')


def cell_snapshot(wb, exclude_run_row=None):
    """Semantic and cell-format preservation check, including blank styled cells."""
    result = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            if ws.title == 'Runs' and row[0].row == exclude_run_row:
                continue
            for c in row:
                result[(ws.title, c.coordinate)] = (
                    c.value, c.data_type, tuple(c._style) if c.has_style else None,
                    c.hyperlink.target if c.hyperlink else None,
                    (c.comment.text, c.comment.author) if c.comment else None,
                )
    return result


def update_registry(path: Path, record: dict, *, append: bool, expected_schema: str):
    csv_path = path.with_suffix('.csv')
    lock = path.with_name('.experiment_registry.lock')
    with exclusive_lock(lock, {'experiment_id': record['experiment_id']}):
        old_xlsx_hash, old_csv_hash = sha256(path), sha256(csv_path)
        wb, header_row, headers, records, _, schema = read_registry(path)
        try:
            if schema != expected_schema or set(record) != set(headers):
                raise RuntimeError('登记表 schema 或当前 run 字段不匹配，禁止写入。')
            check_csv_pair(csv_path, headers, records)
            matches = [(n, r) for n, r in records if r['experiment_id'] == record['experiment_id']]
            if append:
                if matches:
                    raise RuntimeError(f'实验 ID 已存在，拒绝覆盖：{record["experiment_id"]}')
                target = max([header_row] + [n for n, _ in records]) + 1
            else:
                if len(matches) != 1:
                    raise RuntimeError('找不到唯一的当前 run 行，禁止更新。')
                target, old = matches[0]
                if (old['stage'] != 'FormalBaseline' or old['status'] != 'Running'
                        or old['run_directory'] != record['run_directory']):
                    raise RuntimeError('只允许更新当前同一路径的 Running 行，不得改历史结果。')
            if record['stage'] != 'FormalBaseline' or type(record['test_used']) is not bool:
                raise RuntimeError('登记阶段或 test_used 不合法。')
            if record['test_used']:
                if append or record['experiment_id'] != 'Y11-BASE-001-R1' or record['status'] == 'Running':
                    raise RuntimeError('只有本次单 run 的最终更新可以记录 test 使用。')
                evidence = json.loads(record['notes'])
                if evidence.get('test_attempts') != 1 or evidence.get('val') in (None, 'NA'):
                    raise RuntimeError('缺少一次 test 前已完成 val 的证据。')
            if record['experiment_id'] == 'Y11-BASE-001-R1' and record['status'] == 'Completed':
                evidence = json.loads(record['notes'])
                if (record['epochs_completed'] != 100 or record['test_used'] is not True
                        or evidence.get('test') in (None, 'NA')):
                    raise RuntimeError('100 epochs / val / 一次 test 未完整完成，不得登记 Completed。')
            if record['include_in_paper'] != 'Candidate':
                raise RuntimeError('FormalBaseline include_in_paper 必须为 Candidate。')
            if record['status'] not in ('Running', 'Completed', 'Failed'):
                raise RuntimeError('不支持的运行状态；中断使用 Failed 并在 notes 说明。')
            # The chosen empty append row can already be present as a blank row in
            # the XLSX dimensions. Only this one row is authorized to change.
            before = cell_snapshot(wb, target)
            ws = wb['Runs']
            if append:
                for col in range(1, len(headers) + 1):
                    ws.cell(target, col)._style = copy(ws.cell(target - 1, col)._style)
                ws.row_dimensions[target].height = ws.row_dimensions[target - 1].height
            for col, key in enumerate(headers, 1):
                ws.cell(target, col).value = record[key]
            if append:
                for table in ws.tables.values():
                    c1, r1, c2, r2 = range_boundaries(table.ref)
                    if r1 == header_row and c1 == 1 and c2 == len(headers) and r2 == target - 1:
                        table.ref = f'A{header_row}:{get_column_letter(c2)}{target}'
                        if table.autoFilter:
                            table.autoFilter.ref = table.ref
                if ws.auto_filter.ref:
                    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{target}'
            token = datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '_' + uuid.uuid4().hex[:8]
            backup_dir = path.parent / 'backups'
            backup_dir.mkdir(exist_ok=True)
            backups = []
            for original in (path, csv_path):
                backup = backup_dir / f'{original.stem}_{token}{original.suffix}'
                with original.open('rb') as src, backup.open('xb') as dst:
                    shutil.copyfileobj(src, dst)
                backups.append(backup)
            staged_xlsx = path.with_name(f'.registry_{token}.xlsx')
            staged_csv = path.with_name(f'.registry_{token}.csv')
            try:
                wb.save(staged_xlsx)
                check_wb, _, check_headers, check_rows, _, check_schema = read_registry(staged_xlsx)
                try:
                    after = cell_snapshot(check_wb, target)
                    changed = [(key, value, after.get(key)) for key, value in before.items() if after.get(key) != value]
                    if changed:
                        raise RuntimeError(f'暂存工作簿改变了历史单元格/格式，拒绝替换：{changed[:3]}')
                    if check_schema != schema:
                        raise RuntimeError('暂存工作簿改变了字段定义。')
                    with staged_csv.open('x', encoding='utf-8-sig', newline='') as stream:
                        writer = csv.writer(stream)
                        writer.writerow(check_headers)
                        for _, row in check_rows:
                            writer.writerow([csv_value(row[key]) for key in check_headers])
                    check_csv_pair(staged_csv, check_headers, check_rows)
                finally:
                    check_wb.close()
                if sha256(path) != old_xlsx_hash or sha256(csv_path) != old_csv_hash:
                    raise RuntimeError('登记表在写入期间被其他程序改动，停止。')
                os.replace(staged_xlsx, path)
                try:
                    os.replace(staged_csv, csv_path)
                except BaseException:
                    # XLSX has changed but CSV has not: restore the exact original bytes.
                    rollback = path.with_name(f'.rollback_{token}.xlsx')
                    shutil.copyfile(backups[0], rollback)
                    os.replace(rollback, path)
                    raise
            finally:
                staged_xlsx.unlink(missing_ok=True)
                staged_csv.unlink(missing_ok=True)
            return [str(p) for p in backups]
        finally:
            wb.close()
